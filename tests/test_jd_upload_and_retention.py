import io
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw
import docx

from backend.main import app
from backend.database import SessionLocal
from backend.models import User
from backend.auth import create_access_token, get_password_hash
from backend.services.ingestion.parsing import is_supported_jd_file, extract_text_from_file


def get_auth_token():
    db = SessionLocal()
    user = db.query(User).first()
    if not user:
        user = User(email="test_user@truhire.com", name="Test Recruiter", hashed_password=get_password_hash("pass123"))
        db.add(user)
        db.commit()
        db.refresh(user)
    token = create_access_token({"sub": user.id})
    db.close()
    return token


def test_supported_jd_file_types():
    assert is_supported_jd_file("role.pdf") is True
    assert is_supported_jd_file("requirements.docx") is True
    assert is_supported_jd_file("spec.doc") is True
    assert is_supported_jd_file("job.txt") is True
    assert is_supported_jd_file("notes.rtf") is True
    assert is_supported_jd_file("listing.png") is True
    assert is_supported_jd_file("spec.jpg") is True
    assert is_supported_jd_file("image.webp") is True
    assert is_supported_jd_file("malicious.exe") is False
    assert is_supported_jd_file("archive.zip") is False


def test_docx_text_extraction():
    doc = docx.Document()
    doc.add_heading("Senior Backend Engineer", 0)
    doc.add_paragraph("Must have experience with Python, FastAPI, and Postgres.")
    buf = io.BytesIO()
    doc.save(buf)
    extracted = extract_text_from_file("engineer.docx", buf.getvalue())
    assert "Senior Backend Engineer" in extracted
    assert "FastAPI" in extracted


def test_image_ocr_text_extraction():
    img = Image.new("RGB", (400, 100), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((20, 40), "Camunda Developer BPMN", fill=(0, 0, 0))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    extracted = extract_text_from_file("camunda.png", buf.getvalue())
    assert "Camunda" in extracted or "Developer" in extracted or "BPMN" in extracted


def test_jd_upload_api():
    client = TestClient(app)
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    img = Image.new("RGB", (450, 100), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((20, 40), "Staff ML Engineer PyTorch", fill=(0, 0, 0))
    buf = io.BytesIO()
    img.save(buf, format="PNG")

    # Upload is two-step: extract text for review, then create the requirement.
    extracted = client.post(
        "/api/job-descriptions/extract-text",
        headers=headers,
        files={"file": ("ml_staff.png", buf.getvalue(), "image/png")},
    )
    assert extracted.status_code == 200, extracted.text
    ex = extracted.json()
    assert ex["file_url"].startswith("/uploads/job_descriptions/")
    res = client.post("/api/job-descriptions", headers=headers, json={
        "title": "Staff ML Engineer", "jd_text": ex["jd_text"] or "Staff ML Engineer PyTorch",
        "file_url": ex["file_url"], "original_filename": ex["original_filename"],
    })
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["title"] == "Staff ML Engineer"
    assert data["original_filename"] == "ml_staff.png"

    # The stored file is only reachable through the authenticated download route.
    assert client.get(ex["file_url"]).status_code == 404
    assert client.get(f"/api/job-descriptions/{data['id']}/file").status_code == 401
    assert client.get(f"/api/job-descriptions/{data['id']}/file", headers=headers).status_code == 200
    # A client-supplied file reference outside JD storage is rejected.
    bad = client.post("/api/job-descriptions", headers=headers, json={
        "title": "x", "jd_text": "y", "file_url": "/uploads/../truhire.db"})
    assert bad.status_code == 400

    # Cleanup
    db = SessionLocal()
    from backend.models import Position
    pos = db.query(Position).filter(Position.id == data["id"]).first()
    if pos:
        db.delete(pos)
        db.commit()
    db.close()


def test_retention_status_api():
    client = TestClient(app)
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    res = client.get("/api/retention/status", headers=headers)
    assert res.status_code == 200
    data = res.json()
    assert data["data_retention_days"] == 30
    assert "expired_candidates_count" in data
    assert "expired_positions_count" in data


def _fake_llm(title, text):
    return {"location": "Pune, India" if "Pune" in text else "", "role": "Backend Developer" if "Backend Developer" in text else "", "work_mode": ""}


def test_extract_details_from_pasted_text(monkeypatch):
    monkeypatch.setattr("backend.services.llm.ai_service.llm_extract_jd_filters", _fake_llm)
    client = TestClient(app)
    headers = {"Authorization": f"Bearer {get_auth_token()}"}
    text = "Backend Developer needed in Pune, India. 5+ years of experience with Python and FastAPI. Build services for our platform."
    res = client.post("/api/job-descriptions/extract-details", headers=headers, json={"title": "", "jd_text": text})
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["location"] == "Pune, India"
    assert data["detected_title"] == "Backend Developer"
    assert data["experience_years"] == 5
    assert client.post("/api/job-descriptions/extract-details", json={"jd_text": text}).status_code == 401


def test_extract_details_leaves_unstated_fields_blank(monkeypatch):
    monkeypatch.setattr("backend.services.llm.ai_service.llm_extract_jd_filters", lambda t, x: None)
    client = TestClient(app)
    headers = {"Authorization": f"Bearer {get_auth_token()}"}
    text = "We are hiring somebody great to work with our team on exciting problems and grow together."
    data = client.post("/api/job-descriptions/extract-details", headers=headers, json={"title": "Great Hire", "jd_text": text}).json()
    assert data == {"detected_title": None, "location": None, "experience_years": None, "suggested_demand_id": None}
    blank = client.post("/api/job-descriptions/extract-details", headers=headers, json={"jd_text": "   "}).json()
    assert blank == {"detected_title": None, "location": None, "experience_years": None, "suggested_demand_id": None}


def test_extract_details_survives_extractor_failure(monkeypatch):
    def boom(title, text):
        raise RuntimeError("model down")

    monkeypatch.setattr("backend.routers.job_descriptions.extract_jd_filters", boom)
    client = TestClient(app)
    headers = {"Authorization": f"Bearer {get_auth_token()}"}
    res = client.post("/api/job-descriptions/extract-details", headers=headers, json={"jd_text": "Needs 3 years of experience in Java."})
    assert res.status_code == 200 and res.json()["location"] is None and res.json()["experience_years"] == 3


def test_jd_saves_without_role_and_keeps_an_existing_one():
    from backend.models import Position
    client = TestClient(app)
    headers = {"Authorization": f"Bearer {get_auth_token()}"}
    created = client.post("/api/job-descriptions", headers=headers, json={
        "title": "Role Free Title", "jd_text": "Java developer with Spring.", "location": "Pune", "experience_years": 4})
    assert created.status_code == 200, created.text
    jd = created.json()
    assert jd["role"] is None and jd["location"] == "Pune" and jd["experience_years"] == 4

    db = SessionLocal()
    pos = db.query(Position).filter(Position.id == jd["id"]).first()
    pos.role = "Legacy Role"
    db.commit()
    db.close()
    # Editing other fields without sending a role leaves the stored (legacy) role alone.
    updated = client.put(f"/api/job-descriptions/{jd['id']}", headers=headers, json={"location": "Mumbai"})
    assert updated.status_code == 200 and updated.json()["role"] == "Legacy Role"

    db = SessionLocal()
    db.query(Position).filter(Position.id == jd["id"]).delete()
    db.commit()
    db.close()

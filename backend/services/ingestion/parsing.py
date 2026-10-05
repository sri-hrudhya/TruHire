import io
import re
from typing import Dict, Any, List, Optional
import pypdf

# Common tech & general skills dictionary for fast heuristic extraction
COMMON_SKILLS = [
    "Python", "JavaScript", "TypeScript", "React", "Vue", "Angular", "Node.js",
    "FastAPI", "Django", "Flask", "SQL", "PostgreSQL", "MySQL", "MongoDB",
    "Redis", "Docker", "Kubernetes", "AWS", "Azure", "GCP", "Git", "CI/CD",
    "GraphQL", "REST API", "Linux", "Terraform", "Java", "C++", "C#", "Go",
    "Rust", "PHP", "Ruby", "Spring Boot", "Microservices", "Machine Learning",
    "Deep Learning", "NLP", "PyTorch", "TensorFlow", "OpenSearch", "Elasticsearch",
    "Pandas", "NumPy", "Scikit-Learn", "Agile", "Scrum", "DevOps", "Cybersecurity",
    "System Design", "Kafka", "RabbitMQ", "HTML", "CSS", "Tailwind CSS", "Next.js"
]

EMAIL_PATTERN = re.compile(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+')
PHONE_PATTERN = re.compile(r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}')
YEARS_EXP_PATTERN = re.compile(r'(\d+(?:\.\d+)?)\+?\s*(?:years?|yrs?)(?:\s+of)?(?:\s+experience)?', re.IGNORECASE)
EDUCATION_KEYWORDS = ["Bachelor", "Master", "PhD", "B.S.", "M.S.", "B.Tech", "M.Tech", "B.E.", "Degree", "University", "College"]


import os
_ocr_engine = None

PDF_EXTENSIONS = {".pdf"}
DOCX_EXTENSIONS = {".docx"}
DOC_EXTENSIONS = {".doc"}
RTF_EXTENSIONS = {".rtf"}
PLAIN_EXTENSIONS = {".txt", ".md", ".text", ".csv"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tiff", ".tif", ".jfif"}

SUPPORTED_JD_EXTENSIONS = (
    PDF_EXTENSIONS | DOCX_EXTENSIONS | DOC_EXTENSIONS | RTF_EXTENSIONS | PLAIN_EXTENSIONS | IMAGE_EXTENSIONS
)


def is_supported_jd_file(filename: str) -> bool:
    ext = os.path.splitext(filename)[1].lower()
    return ext in SUPPORTED_JD_EXTENSIONS


def get_ocr_engine():
    global _ocr_engine
    if _ocr_engine is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
            _ocr_engine = RapidOCR()
        except Exception as exc:
            print(f"Failed to initialize RapidOCR: {exc}")
            return None
    return _ocr_engine


def extract_text_from_image(file_bytes: bytes) -> str:
    """Extract text from image bytes (JPG, PNG, WEBP, BMP, TIFF, etc.) using RapidOCR."""
    try:
        from PIL import Image
        import numpy as np

        image = Image.open(io.BytesIO(file_bytes))
        if image.mode != "RGB":
            image = image.convert("RGB")

        ocr = get_ocr_engine()
        if not ocr:
            return ""

        result, _ = ocr(np.array(image))
        if not result:
            return ""

        lines = [item[1] for item in result if item and len(item) > 1 and item[1]]
        return "\n".join(lines).strip()
    except Exception as exc:
        print(f"Error in OCR text extraction: {exc}")
        return ""


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from .docx bytes using python-docx with zip/xml fallback."""
    paragraphs = []
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                paragraphs.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        if paragraphs:
            return "\n\n".join(paragraphs).strip()
    except Exception as e:
        print(f"python-docx extraction failed, trying zip/xml fallback: {e}")

    try:
        import zipfile
        import xml.etree.ElementTree as ET
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
            if "word/document.xml" in z.namelist():
                xml_content = z.read("word/document.xml")
                root = ET.fromstring(xml_content)
                texts = [elem.text for elem in root.iter() if elem.text and elem.tag.endswith("t")]
                return " ".join(texts).strip()
    except Exception as e:
        print(f"Fallback docx extraction failed: {e}")
    return ""


def extract_text_from_rtf(file_bytes: bytes) -> str:
    """Extract readable text from RTF content."""
    try:
        raw = file_bytes.decode("latin-1", errors="ignore")
        raw = re.sub(r"\{\*?\\[^{}]+;?\}", "", raw)
        raw = re.sub(r"\\[a-zA-Z]+-?\d*\s?", " ", raw)
        raw = re.sub(r"[{}\\]", "", raw)
        lines = [line.strip() for line in raw.splitlines() if line.strip()]
        return "\n".join(lines).strip()
    except Exception as e:
        print(f"RTF extraction error: {e}")
        return ""


def extract_text_from_doc(file_bytes: bytes) -> str:
    """Extract text from legacy binary Word .doc format."""
    try:
        strings = re.findall(rb"[\x20-\x7E\r\n\t]{4,}", file_bytes)
        decoded = [s.decode("latin-1", errors="ignore").strip() for s in strings if len(s.decode("latin-1", errors="ignore").strip()) > 3]
        return "\n".join(decoded).strip()
    except Exception as e:
        print(f"DOC extraction error: {e}")
        return ""


def extract_text_from_plain(file_bytes: bytes) -> str:
    """Extract text from plain text formats with encoding fallbacks."""
    for enc in ("utf-8", "utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return file_bytes.decode(enc).strip()
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("latin-1", errors="ignore").strip()


def extract_text_from_pdf(pdf_bytes: bytes) -> str:
    """Extract raw text from PDF bytes using pypdf, with OCR fallback for scanned pages."""
    text = ""
    try:
        reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"

        # If very little or no text was found, attempt OCR on page images (scanned PDF)
        if len(text.strip()) < 50 and reader.pages:
            ocr_text_parts = []
            for page in reader.pages:
                try:
                    for img_obj in getattr(page, "images", []):
                        ocr_t = extract_text_from_image(img_obj.data)
                        if ocr_t:
                            ocr_text_parts.append(ocr_t)
                except Exception:
                    pass
            if ocr_text_parts:
                text = "\n\n".join(ocr_text_parts)
    except Exception as e:
        print(f"Error extracting PDF text: {e}")
    return text.strip()


def extract_text_from_file(filename: str, file_bytes: bytes) -> str:
    """
    Extract text based on file format.
    Supports PDF, Word Documents (.docx, .doc), RTF, plain text (.txt, .md),
    and image formats (.jpg, .jpeg, .png, .webp, .bmp, .tiff).
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext in PDF_EXTENSIONS:
        return extract_text_from_pdf(file_bytes)
    elif ext in DOCX_EXTENSIONS:
        return extract_text_from_docx(file_bytes)
    elif ext in DOC_EXTENSIONS:
        return extract_text_from_doc(file_bytes)
    elif ext in RTF_EXTENSIONS:
        return extract_text_from_rtf(file_bytes)
    elif ext in IMAGE_EXTENSIONS:
        return extract_text_from_image(file_bytes)
    elif ext in PLAIN_EXTENSIONS or ext == "":
        return extract_text_from_plain(file_bytes)
    else:
        # Generic fallback: try plain text
        return extract_text_from_plain(file_bytes)


def heuristic_parse_resume(raw_text: str, filename: str) -> Dict[str, Any]:
    """
    Fast regex and keyword extraction of resume fields.
    """
    lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
    
    # 1. Candidate Name
    candidate_name = ""
    if lines:
        # First non-empty line often contains the candidate's name
        first_line = lines[0]
        # Clean obvious header words
        if not any(k in first_line.lower() for k in ["resume", "curriculum", "vitae", "profile", "page"]):
            if len(first_line.split()) <= 4 and len(first_line) < 50:
                candidate_name = first_line

    if not candidate_name:
        # Fallback to filename without extension
        clean_name = filename.rsplit(".", 1)[0].replace("_", " ").replace("-", " ")
        candidate_name = clean_name.title()

    # 2. Email
    emails = EMAIL_PATTERN.findall(raw_text)
    email = emails[0] if emails else None

    # 3. Phone
    phones = PHONE_PATTERN.findall(raw_text)
    phone = phones[0] if phones else None

    # 4. Skills extraction
    found_skills = []
    text_lower = raw_text.lower()
    for skill in COMMON_SKILLS:
        pattern = r'\b' + re.escape(skill.lower()) + r'\b'
        if re.search(pattern, text_lower):
            found_skills.append(skill)

    # 5. Years of experience
    years_exp = 0.0
    matches = YEARS_EXP_PATTERN.findall(raw_text)
    if matches:
        # Take the maximum reasonable stated experience
        exp_nums = []
        for m in matches:
            try:
                val = float(m)
                if 0.5 <= val <= 40.0:
                    exp_nums.append(val)
            except ValueError:
                pass
        if exp_nums:
            years_exp = max(exp_nums)
    else:
        # Estimate based on date ranges (e.g., 2018 - 2023)
        year_matches = re.findall(r'\b(19\d{2}|20\d{2})\s*[-–—to]+\s*(19\d{2}|20\d{2}|present|current)\b', raw_text, re.IGNORECASE)
        if year_matches:
            current_year = 2026
            total_duration = 0.0
            for start, end in year_matches:
                try:
                    s_yr = int(start)
                    e_yr = current_year if end.lower() in ['present', 'current'] else int(end)
                    if 1980 <= s_yr <= e_yr <= current_year:
                        total_duration += (e_yr - s_yr)
                except Exception:
                    pass
            if total_duration > 0:
                years_exp = min(total_duration, 40.0)

    # 6. Education
    education = "Not Specified"
    edu_lines = []
    for line in lines:
        if any(keyword.lower() in line.lower() for keyword in EDUCATION_KEYWORDS):
            edu_lines.append(line)
            if len(edu_lines) >= 2:
                break
    if edu_lines:
        education = " | ".join(edu_lines[:2])

    return {
        "candidate_name": candidate_name,
        "email": email,
        "phone": phone,
        "extracted_skills": found_skills,
        "years_experience": round(years_exp, 1),
        "education": education,
        "raw_text": raw_text
    }


def parse_resume(raw_text: str, filename: str) -> Dict[str, Any]:
    """
    Resume field extraction: the regex heuristic above is the always-available
    baseline (and its skill scan is a fast, deterministic dictionary match), upgraded
    with LLM-based extraction when available. The LLM catches what regex structurally
    can't: free-form skills outside COMMON_SKILLS (e.g. "Camunda", "BPMN"), section
    headers wrongly picked up as names, and experience/education stated in prose
    rather than a fixed pattern. Falls back to the heuristic result untouched if the
    LLM call fails for any reason.
    """
    baseline = heuristic_parse_resume(raw_text, filename)
    try:
        from backend.services.llm.ai_service import llm_parse_resume
        llm_result = llm_parse_resume(raw_text, filename)
    except Exception as exc:
        print(f"LLM resume parse unavailable: {exc}")
        llm_result = None

    if not llm_result:
        return baseline

    merged_skills = list(dict.fromkeys(
        [s for s in llm_result["extracted_skills"] if s] + baseline["extracted_skills"]
    ))

    return {
        "candidate_name": llm_result["candidate_name"] or baseline["candidate_name"],
        "email": baseline["email"],
        "phone": baseline["phone"],
        "extracted_skills": merged_skills,
        "years_experience": llm_result["years_experience"] if llm_result["years_experience"] > 0 else baseline["years_experience"],
        "education": llm_result["education"] if llm_result["education"] != "Not Specified" else baseline["education"],
        "raw_text": raw_text,
    }
